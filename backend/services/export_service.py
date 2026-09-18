from __future__ import annotations

from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from schemas.report import (
    CustomerReportResponse,
    InventoryReportResponse,
    MonthlyBusinessReportResponse,
    ProductReportResponse,
    PurchaseReportResponse,
    SalesReportResponse,
    SupplierReportResponse,
)


class ExportService:
    """Generate Excel and PDF files from existing report data."""

    # ---------------------------------------------------------
    # Excel helpers
    # ---------------------------------------------------------

    @staticmethod
    def _create_workbook(
        title: str,
        headers: list[str],
        rows: Iterable[Iterable[Any]],
    ) -> BytesIO:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = title[:31]

        # Header
        for column_index, header in enumerate(headers, start=1):
            cell = worksheet.cell(
                row=1,
                column=column_index,
                value=header,
            )
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="0F4C5C",
            )
            cell.alignment = Alignment(
                horizontal="center",
            )

        # Data
        for row_index, row_data in enumerate(rows, start=2):
            for column_index, value in enumerate(
                row_data,
                start=1,
            ):
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=value,
                )

        # Auto-size columns
        for column_cells in worksheet.columns:
            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:
                value = (
                    ""
                    if cell.value is None
                    else str(cell.value)
                )

                max_length = max(
                    max_length,
                    len(value),
                )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 12),
                40,
            )

        # Freeze header
        worksheet.freeze_panes = "A2"

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        return output

    # ---------------------------------------------------------
    # Sales Excel
    # ---------------------------------------------------------

    @staticmethod
    def sales_excel(
        report: SalesReportResponse,
    ) -> BytesIO:
        rows = [
            (
                row.invoice_number,
                row.customer_name,
                row.sale_date.replace(tzinfo=None),
                row.payment_method,
                row.status,
                row.subtotal,
                row.discount,
                row.tax,
                row.total,
            )
            for row in report.rows
        ]

        return ExportService._create_workbook(
            title="Sales Report",
            headers=[
                "Invoice Number",
                "Customer",
                "Sale Date",
                "Payment Method",
                "Status",
                "Subtotal",
                "Discount",
                "Tax",
                "Total",
            ],
            rows=rows,
        )

    # ---------------------------------------------------------
    # Inventory Excel
    # ---------------------------------------------------------

    @staticmethod
    def inventory_excel(
        report: InventoryReportResponse,
    ) -> BytesIO:
        rows = [
            (
                row.product_name,
                row.sku,
                row.current_stock,
                row.opening_stock,
                row.damaged_stock,
                row.min_stock_level,
            )
            for row in report.rows
        ]

        return ExportService._create_workbook(
            title="Inventory Report",
            headers=[
                "Product",
                "SKU",
                "Current Stock",
                "Opening Stock",
                "Damaged Stock",
                "Minimum Stock Level",
            ],
            rows=rows,
        )

    # ---------------------------------------------------------
    # Customer Excel
    # ---------------------------------------------------------

    @staticmethod
    def customer_excel(
        report: CustomerReportResponse,
    ) -> BytesIO:
        rows = [
            (
                row.customer_name,
                row.phone,
                row.email,
                row.total_purchases,
                row.total_spending,
            )
            for row in report.rows
        ]

        return ExportService._create_workbook(
            title="Customer Report",
            headers=[
                "Customer",
                "Phone",
                "Email",
                "Total Purchases",
                "Total Spending",
            ],
            rows=rows,
        )

    # ---------------------------------------------------------
    # Supplier Excel
    # ---------------------------------------------------------

    @staticmethod
    def supplier_excel(
        report: SupplierReportResponse,
    ) -> BytesIO:
        rows = [
            (
                row.supplier_name,
                row.total_purchases,
                row.total_purchase_amount,
                row.pending_amount,
            )
            for row in report.rows
        ]

        return ExportService._create_workbook(
            title="Supplier Report",
            headers=[
                "Supplier",
                "Total Purchases",
                "Total Purchase Amount",
                "Pending Amount",
            ],
            rows=rows,
        )

    # ---------------------------------------------------------
    # Product Excel
    # ---------------------------------------------------------

    @staticmethod
    def product_excel(
        report: ProductReportResponse,
    ) -> BytesIO:
        rows = [
            (
                row.product_name,
                row.sku,
                row.quantity_sold,
                row.revenue,
            )
            for row in report.rows
        ]

        return ExportService._create_workbook(
            title="Product Sales",
            headers=[
                "Product",
                "SKU",
                "Quantity Sold",
                "Revenue",
            ],
            rows=rows,
        )

    # ---------------------------------------------------------
    # Purchase Excel
    # ---------------------------------------------------------

    @staticmethod
    def purchase_excel(
        report: PurchaseReportResponse,
    ) -> BytesIO:
        rows = [
            (
                row.supplier_name,
                row.purchase_date.replace(tzinfo=None),
                row.payment_status,
                row.purchase_status,
                row.total_cost,
                row.notes,
            )
            for row in report.rows
        ]

        return ExportService._create_workbook(
            title="Purchase Report",
            headers=[
                "Supplier",
                "Purchase Date",
                "Payment Status",
                "Purchase Status",
                "Total Cost",
                "Notes",
            ],
            rows=rows,
        )

    # ---------------------------------------------------------
    # PDF helpers
    # ---------------------------------------------------------

    @staticmethod
    def _build_pdf(
        title: str,
        headers: list[str],
        rows: list[list[Any]],
        summary: list[tuple[str, Any]] | None = None,
    ) -> BytesIO:
        output = BytesIO()

        document = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=24,
            leftMargin=24,
            topMargin=24,
            bottomMargin=24,
        )

        styles = getSampleStyleSheet()

        elements = [
            Paragraph(
                title,
                styles["Title"],
            ),
            Spacer(1, 12),
        ]

        if summary:
            summary_data = [
                [label, str(value)]
                for label, value in summary
            ]

            summary_table = Table(
                summary_data,
                colWidths=[180, 180],
            )

            summary_table.setStyle(
                TableStyle(
                    [
                        (
                            "BACKGROUND",
                            (0, 0),
                            (0, -1),
                            colors.HexColor(
                                "#0F4C5C"
                            ),
                        ),
                        (
                            "TEXTCOLOR",
                            (0, 0),
                            (0, -1),
                            colors.white,
                        ),
                        (
                            "FONTNAME",
                            (0, 0),
                            (-1, -1),
                            "Helvetica",
                        ),
                        (
                            "GRID",
                            (0, 0),
                            (-1, -1),
                            0.5,
                            colors.grey,
                        ),
                        (
                            "PADDING",
                            (0, 0),
                            (-1, -1),
                            6,
                        ),
                    ]
                )
            )

            elements.extend(
                [
                    summary_table,
                    Spacer(1, 16),
                ]
            )

        table_data = [headers]

        for row in rows:
            table_data.append(
                [
                    str(value)
                    if value is not None
                    else ""
                    for value in row
                ]
            )

        table = Table(
            table_data,
            repeatRows=1,
        )

        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(
                            "#0F4C5C"
                        ),
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white,
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.grey,
                    ),
                    (
                        "ALIGN",
                        (0, 0),
                        (-1, -1),
                        "CENTER",
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "MIDDLE",
                    ),
                    (
                        "PADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                ]
            )
        )

        elements.append(table)

        document.build(elements)

        output.seek(0)

        return output

    # ---------------------------------------------------------
    # Sales PDF
    # ---------------------------------------------------------

    @staticmethod
    def sales_pdf(
        report: SalesReportResponse,
    ) -> BytesIO:
        rows = [
            [
                row.invoice_number,
                row.customer_name,
                row.sale_date.strftime(
                    "%Y-%m-%d %H:%M"
                ),
                row.payment_method,
                row.status,
                row.total,
            ]
            for row in report.rows
        ]

        summary = [
            (
                "Total Sales",
                report.summary.total_sales,
            ),
            (
                "Transactions",
                report.summary.total_transactions,
            ),
            (
                "Average Sale",
                report.summary.average_sale,
            ),
        ]

        return ExportService._build_pdf(
            title="Sales Report",
            headers=[
                "Invoice",
                "Customer",
                "Date",
                "Payment",
                "Status",
                "Total",
            ],
            rows=rows,
            summary=summary,
        )

    # ---------------------------------------------------------
    # Monthly Business Report PDF
    # ---------------------------------------------------------

    @staticmethod
    def monthly_business_report_pdf(
        report: MonthlyBusinessReportResponse,
    ) -> BytesIO:
        summary = report.summary

        summary_rows = [
            (
                "Report Period",
                f"{summary.start_date} to {summary.end_date}",
            ),
            (
                "Total Sales",
                summary.total_sales,
            ),
            (
                "Transactions",
                summary.total_transactions,
            ),
            (
                "Average Sale",
                summary.average_sale,
            ),
            (
                "Revenue",
                summary.revenue,
            ),
            (
                "Cost",
                summary.cost,
            ),
            (
                "Profit",
                summary.profit,
            ),
            (
                "Total Purchases",
                summary.total_purchases,
            ),
            (
                "Total Products",
                summary.total_products,
            ),
            (
                "Low Stock Products",
                summary.low_stock_products,
            ),
        ]

        return ExportService._build_pdf(
            title="Monthly Business Report",
            headers=[
                "Metric",
                "Value",
            ],
            rows=summary_rows,
        )


export_service = ExportService()